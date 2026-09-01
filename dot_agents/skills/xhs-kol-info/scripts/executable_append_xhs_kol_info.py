#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Tuple

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

HEADERS = [
    "序号",
    "达人昵称",
    "达人蒲公英地址",
    "手机号",
    "微信号",
    "博主意向",
    "邀约发起时间",
    "邀约回复时间",
    "原跟进状态",
    "跟进状态更新",
    "读取状态",
    "读取时间",
]


def load_rows(path: Path) -> List[Dict[str, Any]]:
    rows = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(rows, list):
        raise SystemExit("--input must be a JSON array of row objects")
    return [r for r in rows if isinstance(r, dict)]


def ensure_workbook(path: Path):
    if path.exists():
        return load_workbook(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "xhs_kol_info"
    ws.append(HEADERS)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return load_workbook(path)


def canonicalize_headers(ws) -> None:
    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(1, col)
        cell.value = header
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")


def row_dict(ws, row_index: int) -> Dict[str, Any]:
    return {header: ws.cell(row_index, col).value for col, header in enumerate(HEADERS, 1)}


def invite_key(row: Dict[str, Any]) -> Tuple[Any, Any, Any, Any]:
    return (
        row.get("达人昵称"),
        row.get("达人蒲公英地址"),
        row.get("邀约发起时间"),
        row.get("邀约回复时间"),
    )


def last_real_row(ws) -> int:
    name_col = HEADERS.index("达人昵称") + 1
    last = 1
    for row_index in range(2, ws.max_row + 1):
        value = ws.cell(row_index, name_col).value
        if value and str(value).strip():
            last = row_index
    return last


def max_serial(ws) -> int:
    serial = 0
    for row_index in range(2, ws.max_row + 1):
        value = ws.cell(row_index, 1).value
        if isinstance(value, int) and value > serial:
            serial = value
    return serial


def format_sheet(ws) -> None:
    widths = [8, 24, 62, 18, 24, 24, 22, 22, 16, 16, 18, 22]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row_cells in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(HEADERS)):
        for cell in row_cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"


def main() -> int:
    parser = argparse.ArgumentParser(description="Append raw xhs-kol-info JSON rows to an Excel workbook")
    parser.add_argument("--input", type=Path, required=True, help="raw JSON array exported from Chrome/CDP")
    parser.add_argument("--workbook", type=Path, required=True, help="target .xlsx workbook")
    parser.add_argument("--no-backup", action="store_true", help="do not create a timestamped backup before editing")
    args = parser.parse_args()

    rows = load_rows(args.input)
    args.workbook.parent.mkdir(parents=True, exist_ok=True)
    backup_path = None
    if args.workbook.exists() and not args.no_backup:
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        backup_path = args.workbook.with_name(f"{args.workbook.stem}-backup-before-append-{stamp}{args.workbook.suffix}")
        shutil.copy2(args.workbook, backup_path)

    wb = ensure_workbook(args.workbook)
    ws = wb.active
    if ws.title != "xhs_kol_info":
        ws.title = "xhs_kol_info"
    canonicalize_headers(ws)

    last_row = last_real_row(ws)
    removed_ghost_rows = 0
    if ws.max_row > last_row:
        removed_ghost_rows = ws.max_row - last_row
        ws.delete_rows(last_row + 1, removed_ghost_rows)

    existing_keys = set()
    for row_index in range(2, ws.max_row + 1):
        row = row_dict(ws, row_index)
        if row.get("达人昵称"):
            existing_keys.add(invite_key(row))

    serial = max_serial(ws)
    appended = []
    skipped = []
    for row in rows:
        key = invite_key(row)
        if key in existing_keys:
            skipped.append(row)
            continue
        serial += 1
        next_row = dict(row)
        next_row["序号"] = serial
        ws.append([next_row.get(header, "") for header in HEADERS])
        existing_keys.add(key)
        appended.append(next_row)

    format_sheet(ws)
    wb.save(args.workbook)

    print(json.dumps({
        "workbook": str(args.workbook),
        "backup": str(backup_path) if backup_path else "",
        "input": str(args.input),
        "removed_ghost_rows": removed_ghost_rows,
        "appended": len(appended),
        "skipped_exact_duplicates": len(skipped),
        "max_row": ws.max_row,
        "names": [row.get("达人昵称", "") for row in appended],
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
