#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import random
import re
import sys
import time
import urllib.request
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

import websocket  # type: ignore
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

DEFAULT_URL = "https://pgy.xiaohongshu.com/solar/pre-trade/brand/invite-list/note"
SENSITIVE_WARNING = "联系方式属于博主的重要个人信息"
DEFAULT_CDP_PORTS = [9222, 9223, 9224, 9225, 9230]


def human_sleep(min_s: float, max_s: float) -> None:
    time.sleep(random.uniform(min_s, max_s))


class CDP:
    def __init__(self, ws_url: str):
        self.ws = websocket.create_connection(ws_url, timeout=15, suppress_origin=True)
        self.ws.settimeout(15)
        self.seq = 0
        self.call("Runtime.enable")
        self.call("DOM.enable")

    def call(self, method: str, params: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
        self.seq += 1
        self.ws.send(json.dumps({"id": self.seq, "method": method, "params": params or {}}))
        while True:
            msg = json.loads(self.ws.recv())
            if msg.get("id") == self.seq:
                if "error" in msg:
                    raise RuntimeError(f"CDP {method} failed: {msg['error']}")
                return msg

    def eval(self, expression: str, await_promise: bool = True) -> Any:
        res = self.call("Runtime.evaluate", {"expression": expression, "returnByValue": True, "awaitPromise": await_promise})
        result_root = res.get("result", {})
        if "exceptionDetails" in result_root:
            raise RuntimeError(json.dumps(result_root["exceptionDetails"], ensure_ascii=False))
        return result_root.get("result", {}).get("value")

    def click_at(self, x: float, y: float) -> None:
        x += random.uniform(-2.5, 2.5)
        y += random.uniform(-2.5, 2.5)
        self.eval(f"""
(() => {{
  const x = {x};
  const y = {y};
  const el = document.elementFromPoint(x, y);
  if (!el) return false;
  for (const type of ['mouseover', 'mousemove', 'mousedown', 'mouseup', 'click']) {{
    el.dispatchEvent(new MouseEvent(type, {{bubbles:true, cancelable:true, view:window, clientX:x, clientY:y}}));
  }}
  return true;
}})()
""")
        human_sleep(0.08, 0.22)

    def move_to(self, x: float, y: float) -> None:
        x += random.uniform(-2.5, 2.5)
        y += random.uniform(-2.5, 2.5)
        self.eval(f"""
(() => {{
  const x = {x};
  const y = {y};
  const el = document.elementFromPoint(x, y);
  if (!el) return false;
  el.dispatchEvent(new MouseEvent('mouseover', {{bubbles:true, cancelable:true, view:window, clientX:x, clientY:y}}));
  el.dispatchEvent(new MouseEvent('mousemove', {{bubbles:true, cancelable:true, view:window, clientX:x, clientY:y}}));
  return true;
}})()
""")


def browser_display_name(browser: str) -> str:
    return {"chrome": "Chrome", "dia": "Dia", "custom": "浏览器"}.get(browser, browser)


def load_cdp_tabs(port: int) -> List[Dict[str, Any]]:
    import subprocess
    result = subprocess.run(
        ["curl", "--noproxy", "*", "-s", "--max-time", "8", f"http://127.0.0.1:{port}/json/list"],
        capture_output=True, text=True
    )
    if result.returncode != 0 or not result.stdout.strip():
        raise OSError(f"curl failed (rc={result.returncode}): {result.stderr.strip()}")
    return json.loads(result.stdout)


def get_target_ws(port: Optional[int], url: str, browser: str) -> str:
    ports = [port] if port else DEFAULT_CDP_PORTS
    page_snapshots: List[str] = []
    errors: List[str] = []
    label = browser_display_name(browser)
    for candidate in ports:
        if candidate is None:
            continue
        try:
            tabs = load_cdp_tabs(candidate)
        except Exception as exc:
            errors.append(f"- {candidate}: {type(exc).__name__}: {exc}")
            continue
        pages = [t for t in tabs if t.get("type") == "page"]
        target = [t for t in pages if t.get("url") == url] or [
            t for t in pages if "pgy.xiaohongshu.com/solar/pre-trade/brand/invite-list/note" in t.get("url", "")
        ]
        if target:
            print(f"INFO: connected to {label} CDP on 127.0.0.1:{candidate}", file=sys.stderr, flush=True)
            return target[0]["webSocketDebuggerUrl"]
        if pages:
            page_snapshots.append(
                f"127.0.0.1:{candidate}\n"
                + "\n".join(f"- {p.get('title')} | {p.get('url')}" for p in pages[:20])
            )

    if page_snapshots:
        raise SystemExit(
            f"未找到蒲公英邀约列表 {label} 标签页。请先在 {label} 打开并登录目标页面：{url}\n"
            f"当前可见页面：\n" + "\n\n".join(page_snapshots)
        )

    checked = ", ".join(str(p) for p in ports if p is not None)
    detail = "\n".join(errors[:10])
    raise SystemExit(
        f"未找到可用的 {label} CDP 端口。已检查：{checked}。\n"
        f"注意：Computer Use/Chrome 已连接不等于本脚本可访问 Chrome CDP；本脚本需要 Chrome 以 remote debugging 端口暴露页面。\n"
        f"建议：在 Chrome 打开并登录蒲公英目标页后，用 --port <端口> 指向可用 CDP；或用带 --remote-debugging-port=9222 的 Chrome 会话运行。\n"
        f"端口检查结果：\n{detail}"
    )


def parse_value(row_text: Optional[str], label: str) -> str:
    if not row_text:
        return ""
    lines = [ln.strip() for ln in row_text.splitlines() if ln.strip()]
    for i, ln in enumerate(lines):
        if ln == label and i + 1 < len(lines):
            value = lines[i + 1]
            return "" if value.startswith(SENSITIVE_WARNING) else value
    return ""


def extract_invite_name(title_text: Optional[str]) -> str:
    if not title_text:
        return ""
    m = re.search(r"向\s*(.*?)\s*发起合作邀约", title_text)
    return m.group(1).strip() if m else title_text.strip()


def js_string(s: str) -> str:
    return json.dumps(s, ensure_ascii=False)


def check_page_safe(cdp: CDP, browser: str) -> None:
    text = cdp.eval("(() => document.body ? document.body.innerText.slice(0, 3000) : '')()") or ""
    if any(word in text for word in ["验证码", "安全验证", "操作频繁"]) or ("登录" in text and "我的邀约" not in text):
        raise SystemExit(f"页面可能处于登录/验证/风控状态，已停止。请在 {browser_display_name(browser)} 中手动处理后重试。")


def ensure_interested_tab(cdp: CDP, min_delay: float, max_delay: float) -> None:
    state = cdp.eval(r"""
(() => {
  const rows = [...document.querySelectorAll('tr')].filter(tr => tr.innerText.includes('查看详情')).slice(0, 10).map(tr => tr.innerText);
  return {interestedRows: rows.filter(t => t.includes('感兴趣')).length, rowCount: rows.length};
})()
""")
    if state and state.get("rowCount") and state.get("interestedRows") == state.get("rowCount"):
        return
    clicked = cdp.eval(r"""
(() => {
  const el = [...document.querySelectorAll('button,a,span,div')].find(e => /^感兴趣·/.test((e.innerText || '').trim()));
  if (!el) return false;
  el.scrollIntoView({block:'center', inline:'center'});
  el.dispatchEvent(new MouseEvent('mouseover', {bubbles:true, cancelable:true, view:window}));
  el.click();
  return true;
})()
""")
    if clicked:
        human_sleep(min_delay, max_delay)


def apply_interested_unfollowed_filter(cdp: CDP, min_delay: float, max_delay: float) -> None:
    close_drawer(cdp, min_delay * 0.5, max_delay * 0.7)
    ensure_interested_tab(cdp, min_delay, max_delay)
    human_sleep(min_delay * 0.6, max_delay * 0.8)
    opened = cdp.eval(r"""
(() => {
  const labels = [...document.querySelectorAll('span,div,label')];
  const label = labels.find(e => (e.innerText || '').trim() === '跟进状态');
  if (!label) return false;
  let node = label;
  for (let i = 0; i < 6 && node; i++, node = node.parentElement) {
    const candidates = [...node.querySelectorAll('input,[role=combobox],.d-select,.d-select-selector,.d-input-wrapper,button,div')]
      .filter(e => {
        const r = e.getBoundingClientRect();
        const text = (e.innerText || e.getAttribute('placeholder') || e.getAttribute('value') || '').trim();
        return r.width > 40 && r.height > 18 && r.x > label.getBoundingClientRect().x && (text.includes('全部') || text.includes('未跟进') || text.includes('已跟进') || text === '');
      });
    const target = candidates.find(e => {
      const r = e.getBoundingClientRect();
      return r.x > label.getBoundingClientRect().right - 10 && r.x < label.getBoundingClientRect().right + 220;
    });
    if (target) {
      target.scrollIntoView({block:'center', inline:'center'});
      target.dispatchEvent(new MouseEvent('mouseover', {bubbles:true, cancelable:true, view:window}));
      target.click();
      return true;
    }
  }
  return false;
})()
""")
    if opened:
        human_sleep(min_delay * 0.8, max_delay)
        selected = cdp.eval(r"""
(() => {
  const visible = [...document.querySelectorAll('li,div,span')]
    .filter(e => {
      const r = e.getBoundingClientRect();
      const s = getComputedStyle(e);
      return r.width > 0 && r.height > 0 && s.display !== 'none' && s.visibility !== 'hidden';
    });
  const el = visible.find(e => {
    const text = (e.innerText || '').trim();
    return text === '未跟进';
  });
  if (!el) return false;
  el.dispatchEvent(new MouseEvent('mouseover', {bubbles:true, cancelable:true, view:window}));
  el.click();
  return true;
})()
""")
        if selected:
            human_sleep(min_delay * 0.8, max_delay)
    queried = cdp.eval(r"""
(() => {
  const buttons = [...document.querySelectorAll('button,a,[role=button],span,div')]
    .filter(e => (e.innerText || '').trim() === '查询');
  const el = buttons.find(e => {
    const r = e.getBoundingClientRect();
    const s = getComputedStyle(e);
    return r.width > 0 && r.height > 0 && s.display !== 'none' && s.visibility !== 'hidden';
  });
  if (!el) return false;
  el.scrollIntoView({block:'center', inline:'center'});
  el.dispatchEvent(new MouseEvent('mouseover', {bubbles:true, cancelable:true, view:window}));
  el.click();
  return true;
})()
""")
    if queried:
        human_sleep(max(min_delay, 1.5), max_delay + 1.5)


def get_rows(cdp: CDP) -> List[Dict[str, Any]]:
    return cdp.eval(r"""
(() => [...document.querySelectorAll('tr')]
  .map((tr, index) => ({index, text: tr.innerText, hasDetail: [...tr.querySelectorAll('a,button,[role=button]')].some(e => (e.innerText || '').trim() === '查看详情')}))
  .filter(r => r.hasDetail && r.text.includes('感兴趣'))
  .map(r => {
    const lines = r.text.split('\n').map(s => s.trim()).filter(Boolean);
    const skip = new Set(['未跟进','已跟进','感兴趣','喜闻','视频笔记','查看详情','至']);
    const name = lines.find(s => !skip.has(s) && !s.startsWith('¥') && !/^202\d-/.test(s)) || '';
    return {index: r.index, name, text: r.text};
  }))()
""") or []


def get_first_unfollowed_interested_row(cdp: CDP) -> Optional[Dict[str, Any]]:
    rows = get_rows(cdp)
    for row in rows:
        text = row.get("text", "")
        if "未跟进" in text and "感兴趣" in text:
            return row
    return None


def row_key(row: Dict[str, Any]) -> str:
    return f"{row.get('name', '')}|{str(row.get('text', ''))[:500]}"


def get_next_unread_unfollowed_interested_row(cdp: CDP, seen_keys: set[str]) -> Optional[Dict[str, Any]]:
    rows = get_rows(cdp)
    for row in rows:
        text = row.get("text", "")
        key = row_key(row)
        if "未跟进" in text and "感兴趣" in text and key not in seen_keys:
            return row
    return None


def close_drawer(cdp: CDP, min_delay: float, max_delay: float) -> None:
    clicked = cdp.eval(r"""
(() => {
  const scope = document.querySelector('.d-drawer-content,.d-drawer') || document;
  const candidates = [...scope.querySelectorAll('button,a,[role=button],.d-drawer-close,.d-modal-close,span')]
    .filter(e => (e.innerText || '').trim() === '关闭' || /close/i.test(String(e.className)));
  const el = candidates.find(e => { const r=e.getBoundingClientRect(); const s=getComputedStyle(e); return r.width>0 && r.height>0 && s.display !== 'none' && s.visibility !== 'hidden'; });
  if (!el) return false;
  el.dispatchEvent(new MouseEvent('mouseover', {bubbles:true, cancelable:true, view:window}));
  el.click();
  return true;
})()
""")
    if clicked:
        human_sleep(min_delay, max_delay)


def open_detail_by_row_index(cdp: CDP, row_index: int, min_delay: float, max_delay: float) -> None:
    clicked = cdp.eval(f"""
(() => {{
  const trs = [...document.querySelectorAll('tr')];
  const tr = trs[{row_index}];
  if (!tr) return false;
  tr.scrollIntoView({{block:'center', inline:'nearest'}});
  const el = [...tr.querySelectorAll('a,button,[role=button]')].find(e => (e.innerText || '').trim() === '查看详情');
  if (!el) return false;
  el.dispatchEvent(new MouseEvent('mouseover', {{bubbles:true, cancelable:true, view:window}}));
  el.click();
  return true;
}})()
""")
    if not clicked:
        raise RuntimeError(f"找不到第 {row_index} 行的 查看详情")
    human_sleep(min_delay, max_delay)


def mark_row_followed(cdp: CDP, row_index: int, min_delay: float, max_delay: float) -> str:
    selected_row = cdp.eval(f"""
(() => {{
  const trs = [...document.querySelectorAll('tr')];
  const tr = trs[{row_index}];
  if (!tr) return {{ok:false, reason:'row_missing'}};
  tr.scrollIntoView({{block:'center', inline:'nearest'}});
  const input = tr.querySelector('input[type=checkbox]');
  if (input && input.checked) return {{ok:true, rowText:tr.innerText, alreadySelected:true}};
  const checkbox = [...tr.querySelectorAll('.d-grid.d-checkbox.d-checkbox-main.d-clickable')]
    .find(e => {{
      const r = e.getBoundingClientRect();
      return r.width > 0 && r.height > 0;
    }});
  if (!checkbox) return {{ok:false, reason:'checkbox_missing', rowText:tr.innerText}};
  checkbox.dispatchEvent(new MouseEvent('mouseover', {{bubbles:true, cancelable:true, view:window}}));
  checkbox.dispatchEvent(new MouseEvent('mousedown', {{bubbles:true, cancelable:true, view:window}}));
  checkbox.dispatchEvent(new MouseEvent('mouseup', {{bubbles:true, cancelable:true, view:window}}));
  checkbox.click();
  return {{ok:true, rowText:tr.innerText}};
}})()
""")
    if not selected_row or not selected_row.get("ok"):
        return f"SKIP: {selected_row.get('reason') if selected_row else 'checkbox_missing'}"
    human_sleep(min_delay * 0.7, max_delay)
    enabled = cdp.eval(r"""
(() => {
  const btn = [...document.querySelectorAll('button')]
    .find(e => (e.innerText || '').trim() === '标记为已跟进');
  return !!(btn && !btn.disabled && !String(btn.className).includes('disabled'));
})()
""")
    if not enabled:
        return "SKIP: checkbox_not_selected"
    clicked_batch = cdp.eval(r"""
(() => {
  const buttons = [...document.querySelectorAll('button,a,[role=button],span,div')]
    .filter(e => (e.innerText || '').trim() === '标记为已跟进');
  const el = buttons.find(e => e.tagName === 'BUTTON' && !e.disabled && e.getAttribute('aria-disabled') !== 'true' && (() => {
    const r = e.getBoundingClientRect();
    const s = getComputedStyle(e);
    return r.width > 0 && r.height > 0 && s.display !== 'none' && s.visibility !== 'hidden';
  })()) || buttons.find(e => {
    const r = e.getBoundingClientRect();
    const s = getComputedStyle(e);
    return r.width > 0 && r.height > 0 && s.display !== 'none' && s.visibility !== 'hidden' && !String(e.className).includes('disabled');
  });
  if (!el) return false;
  el.dispatchEvent(new MouseEvent('mouseover', {bubbles:true, cancelable:true, view:window}));
  el.click();
  return true;
})()
""")
    if not clicked_batch:
        return "SKIP: batch_follow_button_missing"
    human_sleep(max(min_delay, 1.5), max_delay + 1.5)
    # Some versions open a confirmation popover. Accept the explicit 已跟进/确认 action if present.
    cdp.eval(r"""
(() => {
  const visible = [...document.querySelectorAll('button,a,[role=button],li,div,span')]
    .filter(e => {
      const r = e.getBoundingClientRect();
      const s = getComputedStyle(e);
      return r.width > 0 && r.height > 0 && s.display !== 'none' && s.visibility !== 'hidden';
    });
  const el = visible.find(e => ['确定','确认','已跟进','标记为已跟进'].includes((e.innerText || '').trim()));
  if (el) el.click();
  return true;
})()
""")
    human_sleep(max(min_delay, 1.5), max_delay + 1.5)
    # Refresh current filter so a processed row disappears from 未跟进 results.
    cdp.eval(r"""
(() => {
  const el = [...document.querySelectorAll('button,a,[role=button],span,div')]
    .find(e => (e.innerText || '').trim() === '查询');
  if (el) el.click();
  return true;
})()
""")
    human_sleep(max(min_delay, 1.5), max_delay + 1.5)
    return "已跟进"


def reveal_masked_contacts(cdp: CDP, min_delay: float, max_delay: float) -> None:
    for label in ["博主手机号", "博主微信"]:
        clicked = cdp.eval(f"""
(() => {{
  const drawer = document.querySelector('.d-drawer-content') || document.body;
  const row = [...drawer.querySelectorAll('.flex.gap-16, .detail-item, .detail-data-item, .data-item, .data-list-content')]
    .find(el => (el.innerText || '').includes({js_string(label)}));
  if (!row || !row.innerText.includes('*')) return false;
  const eye = row.querySelector('span.d-icon.cursor-animation, .d-icon.cursor-animation, [class*=\"cursor-animation\"]');
  if (!eye) return false;
  eye.scrollIntoView({{block:'center', inline:'center'}});
  eye.dispatchEvent(new MouseEvent('mouseover', {{bubbles:true, cancelable:true, view:window}}));
  eye.dispatchEvent(new MouseEvent('mousedown', {{bubbles:true, cancelable:true, view:window}}));
  eye.dispatchEvent(new MouseEvent('mouseup', {{bubbles:true, cancelable:true, view:window}}));
  eye.click();
  return true;
}})()
""")
        if clicked:
            human_sleep(0.35, 0.9)
            human_sleep(min_delay * 0.7, max_delay * 0.8)


def read_drawer(cdp: CDP) -> Dict[str, Any]:
    return cdp.eval(r"""
(() => {
  const drawer = document.querySelector('.d-drawer-content') || document.body;
  function rowByLabel(label){return [...drawer.querySelectorAll('.flex.gap-16')].find(el => el.innerText.includes(label));}
  const titleEl = [...drawer.querySelectorAll('a,span,div')].find(e => /^向\s*.+\s*发起合作邀约$/.test((e.innerText || '').trim()));
  const anchors = [...drawer.querySelectorAll('a')].map(a => ({text:a.innerText.trim(), href:a.href}));
  return {
    titleText: titleEl ? titleEl.innerText.trim() : '',
    phoneText: rowByLabel('博主手机号') ? rowByLabel('博主手机号').innerText : '',
    wxText: rowByLabel('博主微信') ? rowByLabel('博主微信').innerText : '',
    intentText: rowByLabel('博主意向') ? rowByLabel('博主意向').innerText : '',
    replyText: rowByLabel('博主回复时间') ? rowByLabel('博主回复时间').innerText : '',
    startText: rowByLabel('邀约发起时间') ? rowByLabel('邀约发起时间').innerText : '',
    anchors
  };
})()
""") or {}


def write_excel(rows: List[Dict[str, Any]], out_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "xhs_kol_info"
    headers = ["序号", "达人昵称", "达人蒲公英地址", "手机号", "微信号", "博主意向", "邀约发起时间", "邀约回复时间", "原跟进状态", "跟进状态更新", "读取状态", "读取时间"]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for i, width in enumerate([8, 24, 62, 18, 24, 24, 22, 22, 16, 16, 18, 22], 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Export Xiaohongshu PGY KOL contact info to Excel")
    parser.add_argument("--count", type=int, required=True, help="number of top interested KOLs to read")
    parser.add_argument("--out", type=Path, default=None, help="output .xlsx path")
    parser.add_argument("--url", default=DEFAULT_URL)
    parser.add_argument("--browser", choices=["chrome", "dia", "custom"], default="chrome", help="browser label for diagnostics; default: chrome")
    parser.add_argument("--port", type=int, default=None, help="Chrome/Dia CDP port; if omitted, tries common local ports")
    parser.add_argument("--min-delay", type=float, default=1.2)
    parser.add_argument("--max-delay", type=float, default=3.4)
    parser.add_argument("--seed", type=int, default=None)
    parser.add_argument("--verbose", action="store_true", help="print step-by-step debug logs")
    parser.add_argument("--mark-followed", action="store_true", help="mark processed rows as 已跟进 after reading (disabled by default)")
    args = parser.parse_args()
    if args.count <= 0:
        raise SystemExit("--count must be positive")
    if args.seed is not None:
        random.seed(args.seed)
    if args.max_delay < args.min_delay:
        raise SystemExit("--max-delay must be >= --min-delay")

    out = args.out or (Path.cwd() / "artifacts" / f"xhs-kol-info-{datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx")
    cdp = CDP(get_target_ws(args.port, args.url, args.browser))
    check_page_safe(cdp, args.browser)
    apply_interested_unfollowed_filter(cdp, args.min_delay, args.max_delay)
    if not get_rows(cdp):
        raise SystemExit("未找到感兴趣 + 未跟进列表行。请确认页面筛选条件和列表已加载。")

    results: List[Dict[str, Any]] = []
    seen_keys: set[str] = set()
    read_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for n in range(1, args.count + 1):
        row = get_first_unfollowed_interested_row(cdp) if args.mark_followed else get_next_unread_unfollowed_interested_row(cdp, seen_keys)
        if not row:
            break
        current_key = row_key(row)
        status = "OK"
        mark_status = "未执行"
        detail: Dict[str, Any] = {}
        try:
            if args.verbose:
                print("DEBUG: before close_drawer", file=sys.stderr, flush=True)
            close_drawer(cdp, args.min_delay * 0.5, args.max_delay * 0.7)
            if args.verbose:
                print("DEBUG: after close_drawer", file=sys.stderr, flush=True)
            human_sleep(args.min_delay, args.max_delay)
            if args.verbose:
                print("DEBUG: before open_detail", row["index"], file=sys.stderr, flush=True)
            open_detail_by_row_index(cdp, int(row["index"]), args.min_delay, args.max_delay)
            if args.verbose:
                print("DEBUG: after open_detail", file=sys.stderr, flush=True)
            reveal_masked_contacts(cdp, args.min_delay, args.max_delay)
            if args.verbose:
                print("DEBUG: after reveal", file=sys.stderr, flush=True)
            detail = read_drawer(cdp)
            if args.verbose:
                print("DEBUG: after read_drawer", file=sys.stderr, flush=True)
        except Exception as exc:
            if args.verbose:
                print("DEBUG: exception", exc, file=sys.stderr, flush=True)
            status = f"ERROR: {exc}"
        title_name = extract_invite_name(detail.get("titleText"))
        anchors = detail.get("anchors") or []
        profile = next((a for a in anchors if a.get("text") == title_name and a.get("href")), None) or next((a for a in anchors if a.get("href") and "/kol/" in a.get("href", "")), None)
        result = {
            "序号": n,
            "达人昵称": row.get("name", ""),
            "达人蒲公英地址": profile.get("href", "") if profile else "",
            "手机号": parse_value(detail.get("phoneText"), "博主手机号"),
            "微信号": parse_value(detail.get("wxText"), "博主微信"),
            "博主意向": parse_value(detail.get("intentText"), "博主意向"),
            "邀约发起时间": parse_value(detail.get("startText"), "邀约发起时间"),
            "邀约回复时间": parse_value(detail.get("replyText"), "博主回复时间"),
            "原跟进状态": "未跟进" if "未跟进" in row.get("text", "") else "",
            "跟进状态更新": mark_status,
            "读取状态": status,
            "读取时间": read_time,
        }
        results.append(result)
        if not args.mark_followed:
            seen_keys.add(current_key)
        close_drawer(cdp, args.min_delay * 0.5, args.max_delay * 0.7)
        if status == "OK" and args.mark_followed:
            try:
                mark_status = mark_row_followed(cdp, int(row["index"]), args.min_delay, args.max_delay)
                result["跟进状态更新"] = mark_status
            except Exception as exc:
                result["跟进状态更新"] = f"ERROR: {exc}"
        human_sleep(args.min_delay, args.max_delay + 1.5)

    write_excel(results, out)
    print(json.dumps({"output": str(out), "count": len(results), "preview": results[:5]}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
