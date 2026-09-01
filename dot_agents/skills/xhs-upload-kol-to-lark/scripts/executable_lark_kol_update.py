#!/usr/bin/env python3
"""Update a Feishu/Lark Base KOL table from local Excel files.

This script intentionally shells out to lark-cli so it inherits the user's
existing Feishu auth, scopes, and safety behavior.
"""

from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
import tempfile
from collections import defaultdict
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlparse

from openpyxl import load_workbook

REQUIRED_FIELDS = ["达人昵称", "手机号", "微信号", "达人蒲公英地址"]
READ_LIMIT = 200
BATCH_CREATE_LIMIT = 500


def clean(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def run_json(args: list[str]) -> dict[str, Any]:
    proc = subprocess.run(args, text=True, capture_output=True)
    if proc.returncode:
        if proc.stdout:
            print(proc.stdout, file=sys.stderr)
        if proc.stderr:
            print(proc.stderr, file=sys.stderr)
        raise SystemExit(proc.returncode)
    try:
        return json.loads(proc.stdout)
    except json.JSONDecodeError as exc:
        print(proc.stdout, file=sys.stderr)
        raise SystemExit(f"Expected JSON from command: {' '.join(args)}") from exc


def parse_target(url: str, base_token: str | None, table_id: str | None) -> tuple[str, str]:
    parsed = urlparse(url)
    query = parse_qs(parsed.query)
    table_from_url = query.get("table", [None])[0]
    table_id = table_id or table_from_url

    if base_token and table_id:
        return base_token, table_id

    wiki_match = re.search(r"/wiki/([^/?#]+)", parsed.path)
    if wiki_match:
        node_token = wiki_match.group(1)
        out = run_json([
            "lark-cli",
            "wiki",
            "spaces",
            "get_node",
            "--params",
            json.dumps({"token": node_token}, ensure_ascii=False),
            "--format",
            "json",
        ])
        node = out.get("data", {}).get("node", {})
        if node.get("obj_type") != "bitable":
            raise SystemExit(f"Wiki node is not a bitable: {node.get('obj_type')}")
        base_token = node.get("obj_token")

    if not base_token:
        base_match = re.search(r"/base/([^/?#]+)", parsed.path)
        if base_match:
            base_token = base_match.group(1)

    if not base_token or not table_id:
        raise SystemExit("Could not resolve base token and table ID from inputs.")
    return base_token, table_id


def read_excel(paths: list[Path]) -> tuple[dict[str, dict[str, str | None]], dict[str, list[tuple[str, int]]], list[dict[str, Any]], int]:
    merged: dict[str, dict[str, str | None]] = {}
    sources: dict[str, list[tuple[str, int]]] = defaultdict(list)
    conflicts: list[dict[str, Any]] = []
    raw_rows = 0

    for path in paths:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            header = [clean(v) or "" for v in next(rows)]
        except StopIteration:
            continue
        index = {name: i for i, name in enumerate(header)}
        missing = [field for field in REQUIRED_FIELDS if field not in index]
        if missing:
            raise SystemExit(f"{path} missing required columns: {missing}")

        for row_num, row in enumerate(rows, start=2):
            nick = clean(row[index["达人昵称"]] if index["达人昵称"] < len(row) else None)
            if not nick:
                continue
            raw_rows += 1
            sources[nick].append((str(path), row_num))
            incoming = {
                field: clean(row[index[field]] if index[field] < len(row) else None)
                for field in REQUIRED_FIELDS
            }
            incoming["达人昵称"] = nick
            if nick not in merged:
                merged[nick] = incoming
                continue
            for field in REQUIRED_FIELDS[1:]:
                old = merged[nick].get(field)
                new = incoming.get(field)
                if new and not old:
                    merged[nick][field] = new
                elif new and old and new != old:
                    conflicts.append({
                        "达人昵称": nick,
                        "字段": field,
                        "existing": old,
                        "incoming": new,
                        "source": str(path),
                        "row": row_num,
                    })
    return merged, sources, conflicts, raw_rows


def field_map(base_token: str, table_id: str) -> dict[str, dict[str, Any]]:
    out = run_json([
        "lark-cli",
        "base",
        "+field-list",
        "--base-token",
        base_token,
        "--table-id",
        table_id,
        "--limit",
        "100",
    ])
    fields = {field["name"]: field for field in out.get("data", {}).get("fields", [])}
    missing = [field for field in REQUIRED_FIELDS if field not in fields]
    if missing:
        raise SystemExit(f"Base table missing required fields: {missing}")
    readonly = {"formula", "lookup", "created_at", "updated_at", "created_by", "updated_by", "auto_number"}
    bad = [name for name in REQUIRED_FIELDS if fields[name].get("type") in readonly]
    if bad:
        raise SystemExit(f"Base fields are not writable: {bad}")
    return fields


def read_existing(base_token: str, table_id: str) -> tuple[dict[str, str], dict[str, list[str]], list[dict[str, Any]]]:
    offset = 0
    by_nick: dict[str, str] = {}
    duplicates: dict[str, list[str]] = defaultdict(list)
    records: list[dict[str, Any]] = []

    while True:
        args = [
            "lark-cli",
            "base",
            "+record-list",
            "--base-token",
            base_token,
            "--table-id",
            table_id,
        ]
        for field in REQUIRED_FIELDS:
            args.extend(["--field-id", field])
        args.extend(["--limit", str(READ_LIMIT), "--offset", str(offset), "--format", "json"])
        out = run_json(args)
        data = out.get("data", {})
        fields = data.get("fields", [])
        nick_index = fields.index("达人昵称")
        for record_id, row in zip(data.get("record_id_list", []), data.get("data", [])):
            record = {"record_id": record_id}
            for field, value in zip(fields, row):
                record[field] = value
            records.append(record)
            nick = clean(row[nick_index])
            if not nick:
                continue
            duplicates[nick].append(record_id)
            if nick not in by_nick:
                by_nick[nick] = record_id
        if not data.get("has_more"):
            break
        offset += READ_LIMIT
    duplicate_map = {nick: ids for nick, ids in duplicates.items() if len(ids) > 1}
    return by_nick, duplicate_map, records


def phone_cell(value: str | None, phone_field_type: str | None) -> int | str | None:
    if not value:
        return None
    if phone_field_type == "number":
        if not value.isdigit():
            raise SystemExit(f"手机号不是纯数字，不能写入数字字段: {value}")
        return int(value)
    return value


def build_payload(record: dict[str, str | None], phone_field_type: str | None) -> dict[str, Any]:
    payload: dict[str, Any] = {"达人昵称": record["达人昵称"]}
    phone = phone_cell(record.get("手机号"), phone_field_type)
    if phone is not None:
        payload["手机号"] = phone
    for field in ["微信号", "达人蒲公英地址"]:
        if record.get(field):
            payload[field] = record[field]
    return payload


def build_batch_row(record: dict[str, str | None], phone_field_type: str | None) -> list[Any]:
    payload = build_payload(record, phone_field_type)
    return [payload.get(field) for field in REQUIRED_FIELDS]


def chunks(values: list[str], size: int) -> list[list[str]]:
    return [values[index:index + size] for index in range(0, len(values), size)]


def write_json_payload(payload: dict[str, Any]) -> Path:
    with tempfile.NamedTemporaryFile(
        "w",
        encoding="utf-8",
        dir=Path.cwd(),
        prefix=".lark_kol_create_",
        suffix=".json",
        delete=False,
    ) as handle:
        json.dump(payload, handle, ensure_ascii=False)
        handle.write("\n")
        return Path(handle.name)


def batch_create(base_token: str, table_id: str, rows: list[list[Any]]) -> None:
    payload_path = write_json_payload({"fields": REQUIRED_FIELDS, "rows": rows})
    try:
        run_json([
            "lark-cli",
            "base",
            "+record-batch-create",
            "--base-token",
            base_token,
            "--table-id",
            table_id,
            "--json",
            f"@./{payload_path.name}",
        ])
    finally:
        payload_path.unlink(missing_ok=True)


def values_match(field: str, expected: str | None, actual: Any) -> bool:
    if not expected:
        return True
    if field == "手机号":
        return clean(actual) == expected
    if field == "达人蒲公英地址" and isinstance(actual, str) and "](" in actual:
        return expected in actual
    return clean(actual) == expected


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--excel", action="append", required=True, type=Path)
    parser.add_argument("--url", required=True)
    parser.add_argument("--base-token")
    parser.add_argument("--table-id")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--write", action="store_true")
    args = parser.parse_args()

    if args.dry_run == args.write:
        raise SystemExit("Pass exactly one of --dry-run or --write.")

    merged, sources, conflicts, raw_rows = read_excel(args.excel)
    duplicate_excel = {nick: locs for nick, locs in sources.items() if len(locs) > 1}
    if conflicts:
        print(json.dumps({"ok": False, "conflicts": conflicts}, ensure_ascii=False, indent=2))
        raise SystemExit(2)

    base_token, table_id = parse_target(args.url, args.base_token, args.table_id)
    fields = field_map(base_token, table_id)
    existing, duplicate_base, _records = read_existing(base_token, table_id)
    if duplicate_base:
        print(json.dumps({"ok": False, "duplicate_base_nicknames": duplicate_base}, ensure_ascii=False, indent=2))
        raise SystemExit(3)

    to_update = [nick for nick in merged if nick in existing]
    to_create = [nick for nick in merged if nick not in existing]
    summary = {
        "ok": True,
        "base_token": base_token,
        "table_id": table_id,
        "excel_raw_rows": raw_rows,
        "unique_nicknames": len(merged),
        "duplicate_excel_nicknames": duplicate_excel,
        "to_update": to_update,
        "to_create": to_create,
    }

    if args.dry_run:
        print(json.dumps(summary, ensure_ascii=False, indent=2))
        return

    phone_type = fields["手机号"].get("type")
    updated: list[str] = []
    for nick in to_update:
        record = merged[nick]
        payload = build_payload(record, phone_type)
        command = [
            "lark-cli",
            "base",
            "+record-upsert",
            "--base-token",
            base_token,
            "--table-id",
            table_id,
            "--json",
            json.dumps(payload, ensure_ascii=False),
        ]
        command[5:5] = ["--record-id", existing[nick]]
        run_json(command)
        updated.append(nick)

    created: list[str] = []
    for nick_chunk in chunks(to_create, BATCH_CREATE_LIMIT):
        rows = [build_batch_row(merged[nick], phone_type) for nick in nick_chunk]
        batch_create(base_token, table_id, rows)
        created.extend(nick_chunk)

    _existing_after, duplicate_after, records_after = read_existing(base_token, table_id)
    if duplicate_after:
        raise SystemExit(f"Duplicate nicknames appeared after write: {duplicate_after}")
    by_nick_after = {record.get("达人昵称"): record for record in records_after}
    missing = [nick for nick in merged if nick not in by_nick_after]
    issues = []
    for nick, expected in merged.items():
        actual_record = by_nick_after.get(nick)
        if not actual_record:
            continue
        for field in REQUIRED_FIELDS[1:]:
            if not values_match(field, expected.get(field), actual_record.get(field)):
                issues.append({
                    "达人昵称": nick,
                    "字段": field,
                    "expected": expected.get(field),
                    "actual": actual_record.get(field),
                })

    print(json.dumps({
        **summary,
        "created": created,
        "updated": updated,
        "verify": {
            "checked": len(merged),
            "found": len(merged) - len(missing),
            "missing": missing,
            "issues": issues,
        },
    }, ensure_ascii=False, indent=2))
    if missing or issues:
        raise SystemExit(4)


if __name__ == "__main__":
    main()
